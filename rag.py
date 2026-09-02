import io
import json
import os
import urllib.request

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

CHUNK_SIZE = 800
CHUNK_OVERLAP = 100

OFFICE_CONTENT_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def guess_content_type(filename):
    lower = filename.lower()
    for ext, content_type in OFFICE_CONTENT_TYPES.items():
        if lower.endswith(ext):
            return content_type
    return "text/plain"


def strip_nul_bytes(text):
    # PostgreSQL TEXT columns reject the NUL (0x00) codepoint. pypdf and other
    # extractors sometimes surface it from malformed source streams.
    return text.replace("\x00", "")


def extract_pdf_text(raw_bytes):
    reader = PdfReader(io.BytesIO(raw_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_docx_text(raw_bytes):
    document = DocxDocument(io.BytesIO(raw_bytes))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def extract_pptx_text(raw_bytes):
    presentation = Presentation(io.BytesIO(raw_bytes))
    parts = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            parts.append(cell.text)
    return "\n".join(parts)


def extract_xlsx_text(raw_bytes):
    workbook = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        parts.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(value) for value in row if value is not None and str(value).strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(filename, raw_bytes):
    lower = filename.lower()
    if lower.endswith(".pdf"):
        text = extract_pdf_text(raw_bytes)
    elif lower.endswith(".docx"):
        text = extract_docx_text(raw_bytes)
    elif lower.endswith(".pptx"):
        text = extract_pptx_text(raw_bytes)
    elif lower.endswith(".xlsx"):
        text = extract_xlsx_text(raw_bytes)
    else:
        text = raw_bytes.decode("utf-8", errors="replace")
    return strip_nul_bytes(text)


def chunk_text(text, size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    text = strip_nul_bytes(text).strip()
    chunks = []
    start = 0
    while start < len(text):
        end = start + size
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = end - overlap
    return chunks


def embed(text, ollama_url=None, model=None):
    ollama_url = (ollama_url or os.getenv("OLLAMA_URL", "http://127.0.0.1:11434")).rstrip("/")
    model = model or os.getenv("EMBEDDING_MODEL", "nomic-embed-text")
    payload = json.dumps({"model": model, "input": text}).encode("utf-8")
    request = urllib.request.Request(
        f"{ollama_url}/api/embed",
        data=payload,
        headers={"Content-Type": "application/json"},
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        data = json.loads(response.read().decode("utf-8"))
    return data["embeddings"][0]


def vector_literal(values):
    return "[" + ",".join(repr(float(v)) for v in values) + "]"


def ingest_document(pool, uploaded_by, filename, content_type, raw_bytes):
    text = extract_text(filename, raw_bytes)
    chunks = chunk_text(text)
    with pool.connection() as conn:
        document_id = conn.execute(
            "INSERT INTO documents (uploaded_by, filename, content_type) "
            "VALUES (%s, %s, %s) RETURNING id",
            (uploaded_by, filename, content_type),
        ).fetchone()[0]
        for index, chunk in enumerate(chunks):
            vector = embed(chunk)
            conn.execute(
                "INSERT INTO document_chunks (document_id, chunk_index, content, embedding) "
                "VALUES (%s, %s, %s, %s)",
                (document_id, index, chunk, vector_literal(vector)),
            )
    return document_id, len(chunks)


def list_documents(pool):
    with pool.connection() as conn:
        rows = conn.execute(
            "SELECT d.id, d.filename, d.content_type, d.created_at, COUNT(c.id) "
            "FROM documents d LEFT JOIN document_chunks c ON c.document_id = d.id "
            "GROUP BY d.id ORDER BY d.created_at DESC"
        ).fetchall()
    return [
        {
            "id": row[0],
            "filename": row[1],
            "content_type": row[2],
            "created_at": row[3].isoformat(),
            "chunks": row[4],
        }
        for row in rows
    ]


def delete_document(pool, document_id):
    with pool.connection() as conn:
        conn.execute("DELETE FROM documents WHERE id = %s", (document_id,))


def retrieve_context(pool, query, top_k=None):
    top_k = top_k or int(os.getenv("RAG_TOP_K", "4"))
    try:
        vector = embed(query)
    except Exception:
        return []
    with pool.connection() as conn:
        rows = conn.execute(
            "SELECT content, embedding <=> %s AS distance FROM document_chunks "
            "ORDER BY distance ASC LIMIT %s",
            (vector_literal(vector), top_k),
        ).fetchall()
    return [row[0] for row in rows]
